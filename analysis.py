import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

HUMAN_FASTA_PATH = "data/human_HBB.fasta"
CHIMP_FASTA_PATH = "data/chimp_HBB.fasta"
PNG_OUT_PATH = "results/figures/hbb_comparison.png"
XLSX_OUT_PATH = "results/hbb_comparison.xlsx"


def read_fasta(path):
    header = None
    sequence_pieces = []

    with open(path, "r") as fasta_file:
        for line in fasta_file:
            line = line.strip()

            if line.startswith(">"):
                header = line[1:]
            else:
                sequence_pieces.append(line)

    sequence = "".join(sequence_pieces)
    return header, sequence


def sequence_length(sequence):
    return len(sequence)


def gc_content(sequence):
    g_count = sequence.count("G")
    c_count = sequence.count("C")
    return 100 * (g_count + c_count) / len(sequence)


def nucleotide_composition(sequence):
    length = len(sequence)
    composition = {}
    for base in "ACGT":
        count = sequence.count(base)
        percentage = 100 * count / length
        composition[base] = (count, percentage)
    return composition


def analyze_human():
    return _analyze(HUMAN_FASTA_PATH)


def analyze_chimp():
    return _analyze(CHIMP_FASTA_PATH)


def _analyze(path):
    header, sequence = read_fasta(path)
    return {
        "header": header,
        "sequence": sequence,
        "length": sequence_length(sequence),
        "gc_content": gc_content(sequence),
        "composition": nucleotide_composition(sequence),
    }


def print_stats(species_name, stats):
    print(f"--- {species_name} ---")
    print(f"Sequence ID / description: {stats['header']}")
    print(f"Sequence length: {stats['length']} bp")
    print(f"GC content: {stats['gc_content']:.2f}%")
    print("Nucleotide composition:")
    for base, (count, percentage) in stats["composition"].items():
        print(f"  {base}: {count:>4} bases ({percentage:.1f}%)")
    print(f"First 60 bases: {stats['sequence'][:60]}")
    print()


def make_png_chart(human_stats, chimp_stats, out_path=PNG_OUT_PATH):
    fig, (ax_gc, ax_comp) = plt.subplots(1, 2, figsize=(10, 4))

    species = ["Human", "Chimpanzee"]
    gc_values = [human_stats["gc_content"], chimp_stats["gc_content"]]
    ax_gc.bar(species, gc_values, color=["#3b7ddd", "#f4a63a"])
    ax_gc.set_ylabel("GC content (%)")
    ax_gc.set_title("GC content")
    ax_gc.set_ylim(0, 100)
    for i, v in enumerate(gc_values):
        ax_gc.text(i, v + 1, f"{v:.1f}%", ha="center")

    bases = ["A", "C", "G", "T"]
    human_counts = [human_stats["composition"][b][0] for b in bases]
    chimp_counts = [chimp_stats["composition"][b][0] for b in bases]

    x = range(len(bases))
    width = 0.35
    ax_comp.bar([i - width / 2 for i in x], human_counts, width, label="Human", color="#3b7ddd")
    ax_comp.bar([i + width / 2 for i in x], chimp_counts, width, label="Chimpanzee", color="#f4a63a")
    ax_comp.set_xticks(list(x))
    ax_comp.set_xticklabels(bases)
    ax_comp.set_ylabel("Base count")
    ax_comp.set_title("Nucleotide composition")
    ax_comp.legend()

    fig.suptitle("HBB gene: Human vs. Chimpanzee")
    fig.tight_layout()
    fig.savefig(out_path, dpi=150)
    plt.close(fig)
    print(f"Chart image saved to {out_path}")


def make_excel_report(human_stats, chimp_stats, out_path=XLSX_OUT_PATH):
    wb = Workbook()
    ws = wb.active
    ws.title = "HBB Comparison"

    default_font = Font(name="Arial", size=11)
    bold_font = Font(name="Arial", size=11, bold=True)

    headers = ["Species", "Length (bp)", "A", "C", "G", "T", "GC Content (%)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = bold_font

    def row_for(stats):
        comp = stats["composition"]
        return [comp["A"][0], comp["C"][0], comp["G"][0], comp["T"][0]]

    ws.append(["Human", human_stats["length"]] + row_for(human_stats) + [None])
    ws.append(["Chimpanzee", chimp_stats["length"]] + row_for(chimp_stats) + [None])

    ws["G2"] = "=(D2+E2)/B2*100"
    ws["G3"] = "=(D3+E3)/B3*100"
    ws["G2"].number_format = "0.0"
    ws["G3"].number_format = "0.0"

    for row in ws.iter_rows(min_row=2, max_row=3, min_col=1, max_col=7):
        for cell in row:
            cell.font = default_font

    ws["A5"] = "Base"
    ws["B5"] = "Human"
    ws["C5"] = "Chimpanzee"
    for cell in ws[5]:
        cell.font = bold_font

    base_to_col = {"A": "C", "C": "D", "G": "E", "T": "F"}
    for i, base in enumerate(["A", "C", "G", "T"]):
        r = 6 + i
        col = base_to_col[base]
        ws[f"A{r}"] = base
        ws[f"B{r}"] = f"={col}2"
        ws[f"C{r}"] = f"={col}3"
        for cell in ws[r]:
            cell.font = default_font

    gc_chart = BarChart()
    gc_chart.title = "GC Content: Human vs. Chimpanzee"
    gc_chart.y_axis.title = "GC content (%)"
    gc_chart.style = 10
    data = Reference(ws, min_col=7, min_row=1, max_row=3)
    categories = Reference(ws, min_col=1, min_row=2, max_row=3)
    gc_chart.add_data(data, titles_from_data=True)
    gc_chart.set_categories(categories)
    ws.add_chart(gc_chart, "I2")

    comp_chart = BarChart()
    comp_chart.type = "col"
    comp_chart.grouping = "clustered"
    comp_chart.title = "Nucleotide Composition"
    comp_chart.y_axis.title = "Base count"
    comp_chart.style = 10
    data = Reference(ws, min_col=2, max_col=3, min_row=5, max_row=9)
    categories = Reference(ws, min_col=1, min_row=6, max_row=9)
    comp_chart.add_data(data, titles_from_data=True)
    comp_chart.set_categories(categories)
    ws.add_chart(comp_chart, "I18")

    for col, width in zip("ABCDEFG", [12, 12, 8, 8, 8, 8, 16]):
        ws.column_dimensions[col].width = width

    wb.save(out_path)
    print(f"Excel workbook saved to {out_path}")


def main():
    human_stats = analyze_human()
    chimp_stats = analyze_chimp()

    print_stats("Human HBB", human_stats)
    print_stats("Chimpanzee HBB", chimp_stats)

    make_png_chart(human_stats, chimp_stats)
    make_excel_report(human_stats, chimp_stats)


if __name__ == "__main__":
    main()
